"""Reportes de disponibilidad en Excel, CSV y PDF (item #7, CU04, RF05).

Una fila por analisis de una gondola: fecha, ubicacion, % libre de la
gondola completa y alerta si/no (el contenido minimo de RF05), mas el
estado de la alerta y quien subio la foto. No se informa nada por repisa
(#40). Recibe filas simples (dicts con fecha, pasillo, gondola,
free_fraction, critical, estado y usuario), sin depender de la app web.
"""

from __future__ import annotations

import csv
import io

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

COLUMNS = ["Fecha", "Pasillo", "Góndola", "% libre", "Alerta", "Estado", "Subida por"]
MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "pdf": "application/pdf",
}


def export(rows: list[dict], fmt: str, title: str) -> bytes:
    """Bytes del reporte en `fmt` ("xlsx", "csv" o "pdf"). `title` va como
    encabezado del PDF y nombre de la hoja de Excel."""
    table = [
        [row["fecha"].strftime("%Y-%m-%d %H:%M"), row["pasillo"], row["gondola"],
         round(row["free_fraction"] * 100, 1), "Sí" if row["critical"] else "No",
         row["estado"], row["usuario"]]
        for row in rows
    ]
    if fmt == "csv":
        return _csv(table)
    if fmt == "xlsx":
        return _xlsx(table, title)
    if fmt == "pdf":
        return _pdf(table, title)
    raise ValueError(f"Formato desconocido: {fmt}")


def _csv(table: list[list]) -> bytes:
    text = io.StringIO()
    csv.writer(text).writerows([COLUMNS, *table])
    # Con BOM, Excel abre el CSV en UTF-8 y muestra bien "Góndola" y "Sí".
    return text.getvalue().encode("utf-8-sig")


def _xlsx(table: list[list], title: str) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]  # limite de Excel para el nombre de una hoja
    sheet.append(COLUMNS)
    for row in table:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFG", (17, 10, 10, 9, 8, 14, 14)):
        sheet.column_dimensions[column].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _pdf(table: list[list], title: str) -> bytes:
    pdf = FPDF(orientation="landscape")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(text=title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    alerts = sum(row[4] == "Sí" for row in table)
    pdf.cell(text=f"{len(table)} análisis, {alerts} con quiebre de stock (alerta de la góndola completa).",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    with pdf.table(col_widths=(34, 18, 18, 16, 14, 40, 30), text_align="LEFT") as pdf_table:
        for row in [COLUMNS, *table]:
            pdf_row = pdf_table.row()
            for value in row:
                pdf_row.cell(str(value))
    return bytes(pdf.output())
