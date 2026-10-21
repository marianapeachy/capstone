"""Reportes Excel, CSV y PDF (item #7, CU04): una fila por analisis con el
% libre de la gondola completa y la alerta si/no."""

import csv
import io
from datetime import datetime

import openpyxl
import pytest

from src.reporting.export import COLUMNS, export

ROWS = [
    {"fecha": datetime(2026, 10, 21, 10, 30), "pasillo": "3", "gondola": "14B",
     "free_fraction": 0.4234, "critical": True, "estado": "Activa", "usuario": "12345678-5"},
    {"fecha": datetime(2026, 10, 21, 11, 5), "pasillo": "3", "gondola": "15",
     "free_fraction": 0.05, "critical": False, "estado": "Sin quiebre", "usuario": "12345678-5"},
]
EXPECTED = [
    ["2026-10-21 10:30", "3", "14B", 42.3, "Sí", "Activa", "12345678-5"],
    ["2026-10-21 11:05", "3", "15", 5.0, "No", "Sin quiebre", "12345678-5"],
]


def test_csv_has_header_and_one_row_per_analysis():
    text = export(ROWS, "csv", "Reporte").decode("utf-8-sig")
    header, *rows = list(csv.reader(io.StringIO(text)))
    assert header == COLUMNS
    assert rows == [[str(v) for v in row] for row in EXPECTED]


def test_csv_starts_with_bom_so_excel_reads_accents():
    assert export(ROWS, "csv", "Reporte").startswith("﻿".encode())


def test_xlsx_keeps_free_space_as_a_number():
    sheet = openpyxl.load_workbook(io.BytesIO(export(ROWS, "xlsx", "Reporte"))).active
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert values == [COLUMNS, *EXPECTED]


def test_pdf_is_a_complete_pdf_file():
    data = export(ROWS, "pdf", "Reporte de disponibilidad")
    assert data.startswith(b"%PDF") and data.rstrip().endswith(b"%%EOF")


def test_unknown_format_is_rejected():
    with pytest.raises(ValueError):
        export(ROWS, "docx", "Reporte")


def test_text_that_looks_like_a_formula_is_not_run_by_excel():
    # Pasillo y góndola los escribe un usuario: "=..." no debe llegar como fórmula.
    rows = [{**ROWS[0], "pasillo": "=1+1", "gondola": "@SUM(A1)"}]
    cell = openpyxl.load_workbook(io.BytesIO(export(rows, "xlsx", "Reporte"))).active["B2"]
    assert cell.data_type == "s" and cell.value == "'=1+1"
    line = export(rows, "csv", "Reporte").decode("utf-8-sig").splitlines()[1]
    assert ",'=1+1,'@SUM(A1)," in line
